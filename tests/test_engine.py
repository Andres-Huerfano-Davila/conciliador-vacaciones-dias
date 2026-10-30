import io
import unittest
from datetime import date
import pandas as pd
from openpyxl import load_workbook
from engine import read_file, reconcile, commercial, excel_bytes, absence_total, dates

def fixtures():
    return {
        'master':pd.DataFrame({'SAP':['1','2'],'Número de personal':['Ana','Luis'],'Área de nómina':['ZM','ZL'],'Área de personal':['Non Manager','Non Manager'],'Fecha':['01.01.2026','01.01.2026'],'Ce.coste':['103','101']}),
        'aus':pd.DataFrame({'SAP':['1','2'],'Cl.':['0100','0100'],'Inicio':['01.02.2026']*2,'Final':['02.02.2026']*2,'Días':['1,00']*2}),
        'book':pd.DataFrame({'SAP':['1','2'],'Cantidad':['15,00']*2,'LiquidCont':['10,00']*2}),
        'prov':pd.DataFrame({'SAP':['1','2'],'TipoCalc':['PVVA']*2,'Dias N Sld':['6,50','6,50'],'Dias Sd A':['5,25','5,25']}),
        'sf':pd.DataFrame({'SAP':['1','2'],'Amount posted':['5.25','5.25'],'Posting Date':['30/06/2026']*2,'Time Account Type External Code':['CO_vacatio']*2}),
        'pending':pd.DataFrame({'SAP':['1'],'approvalStatus':['PENDING'],'quantityInDays':['2'],'externalCode':['CO_vacatio']})}

class Tests(unittest.TestCase):
    def test_three_absence_sources_route_and_deduplicate(self):
        f=fixtures()
        f['master']['Área de nómina']=['ZM','ZH']
        f['aus_full']=f.pop('aus')
        f['aus_pt']=pd.DataFrame({'SAP':['2'], 'Cl.':['0100'], 'Inicio':['01.02.2026'], 'Final':['02.02.2026'], 'Tot.días nóm.':['2,00']})
        f['aus_pt_current']=f['aus_pt'].copy()
        r,audit=reconcile(f,date(2026,6,30),True)
        r=r.set_index('SAP')
        self.assertEqual(r.loc['1','Días disfrutados'],1)
        self.assertEqual(r.loc['2','Días disfrutados'],2)
        self.assertEqual(audit['Duplicados de ausencias eliminados'],1)
        f['aus_pt_current']['Tot.días nóm.']=['3,00']
        r,_=reconcile(f,date(2026,6,30),True)
        self.assertTrue(pd.isna(r.set_index('SAP').loc['2','Cálculo']))
        self.assertFalse(pd.isna(r.set_index('SAP').loc['1','Cálculo']))
        del f['aus_pt_current']
        r,_=reconcile(f,date(2026,6,30),True)
        self.assertEqual(r.set_index('SAP').loc['2','Días disfrutados'],2)
        del f['aus_pt']
        with self.assertRaisesRegex(ValueError,'Faltan archivos'):
            reconcile(f,date(2026,6,30),True)

    def test_user_tolerances_decide_review(self):
        f=fixtures()
        r,_=reconcile(f,date(2026,6,30),True,tolerances={'SF vs SAP 2006':1000})
        self.assertEqual(r['SF vs SAP 2006'].iloc[0],'OK')
        r,audit=reconcile(f,date(2026,6,30),True,tolerances={'SF vs SAP 2006':0})
        self.assertEqual(audit['Tolerancias']['SF vs SAP 2006'],0)
        self.assertEqual(r['SF vs SAP 2006'].iloc[0],'Revisar' if r['Δ SF vs SAP 2006'].iloc[0]!=0 else 'OK')

    def test_reader_retains_part_time_payroll_days(self):
        data='|Nº pers.|ClAbPre|Inicio|Final|Tot.días nóm.|\n|2|0100|01.02.2026|02.02.2026|2,00|\n'
        frame=read_file(data.encode('cp1252'),'AUS_PT_2026.txt')
        self.assertEqual(frame['Tot.días nóm.'].iloc[0],'2,00')

    def test_sf_dates_with_unpadded_time_preserve_day_first(self):
        d=dates(pd.Series(['05/10/2026 0:00:00','30.09.2026','2026-09-30','31.12.9999','']))
        self.assertEqual(d.iloc[0].date(),date(2026,10,5))
        self.assertEqual(d.iloc[1].date(),date(2026,9,30))
        self.assertEqual(d.iloc[2].date(),date(2026,9,30))
        self.assertTrue(d.iloc[3:].isna().all())
    def test_sena_uses_personnel_area_even_when_payroll_zl(self):
        f=fixtures(); f['master']['Área de personal']=['SENA PRODUCTIVO','Non Manager']
        f['master']['Área de nómina']=['MENSUAL ADMON 365','MENSUAL ADMON 365']
        r,_=reconcile(f,date(2026,6,30),True); r=r.set_index('SAP')
        self.assertEqual(r.loc['1','SENA Productivo'],'Sí')
        self.assertEqual(r.loc['1','Días causados'],7.5)
        self.assertAlmostEqual(r.loc['2','Días causados'],181*15/365)
    def test_sena_day31_only_short_original_absence(self):
        def rows(a,b):
            return pd.DataFrame({'code':['0380'],'start':pd.to_datetime([a]),'end':pd.to_datetime([b]),'days':[1.]})
        start,end=date(2026,7,1),date(2026,7,31)
        self.assertEqual(absence_total(rows('2026-07-31','2026-07-31'),start,end,'ZM',{'0380'},sena=True),2)
        self.assertEqual(absence_total(rows('2026-07-31','2026-07-31'),start,end,'ZM',{'0380'},sena=False),1)
        self.assertEqual(absence_total(rows('2026-07-01','2026-08-15'),start,end,'ZM',{'0380'},sena=True),30)
        self.assertEqual(absence_total(rows('2026-08-02','2026-08-31'),date(2026,8,1),date(2026,8,31),'ZM',{'0380'},sena=True),30)
        self.assertEqual(absence_total(rows('2026-08-01','2026-08-31'),date(2026,8,1),date(2026,8,31),'ZM',{'0380'},sena=True),30)
    def test_rehire_and_22_day_entitlement(self):
        f=fixtures(); f['master']['Fecha reingreso']=['01.06.2026','']
        f['master']['Área de nómina']=['ADMINISTRATIVOS VACACIONES 22','ZL']
        r,_=reconcile(f,date(2026,6,30),True); r=r.set_index('SAP')
        self.assertAlmostEqual(r.loc['1','Días causados'],30*22/360)
        self.assertEqual(r.loc['1','Días disfrutados'],0)
    def test_commercial_february(self):
        self.assertEqual(commercial(date(2026,2,1),date(2026,2,28)),30)
    def test_historical_area_divisors_and_consumed_quota(self):
        result,_=reconcile(fixtures(),date(2026,6,30),True)
        r=result.set_index('SAP')
        self.assertAlmostEqual(r.loc['1','Días causados'],7.5)
        self.assertAlmostEqual(r.loc['2','Días causados'],181*15/365)
        self.assertEqual(r.loc['1','SAP 2006'],5)
        self.assertEqual(r.loc['1','Cálculo'],6.5)
        self.assertEqual(r.loc['1','Cálculo vs Provisión'],'OK')
        self.assertEqual(r.loc['1','Pendientes SF'],2)
        self.assertEqual(r.loc['1','SF vs Provisión'],'OK')
        self.assertEqual(r.loc['1','SAP 2006 vs Provisión'],'Revisar')
    def test_missing_and_source_only_not_zero(self):
        f=fixtures(); f['sf']=f['sf'].iloc[:1].copy(); f['sf'].loc[1]=['3','10','30/06/2026','CO_vacatio']
        r,_=reconcile(f,date(2026,6,30),True,include_orphans=True); r=r.set_index('SAP')
        self.assertTrue(pd.isna(r.loc['2','SF']))
        self.assertEqual(r.loc['2','SF vs Provisión'],'Sin información')
        self.assertIn('3',r.index)
        self.assertEqual(r.loc['3','Estado'],'Sin información')
    def test_history_unconfirmed(self):
        r,_=reconcile(fixtures(),date(2026,6,30))
        self.assertTrue(r['Cálculo'].isna().all())
    def test_exact_threshold(self):
        f=fixtures(); f['prov']['Dias N Sld']=['6,25','6,75']
        r,_=reconcile(f,date(2026,6,30),True)
        self.assertEqual(r.set_index('SAP').loc['1','Cálculo vs Provisión'],'OK')
        f['prov']['Dias N Sld']=['6,249','6,75']
        r,_=reconcile(f,date(2026,6,30),True)
        self.assertEqual(r.set_index('SAP').loc['1','Cálculo vs Provisión'],'Revisar')
    def test_unpaid_overlap(self):
        f=fixtures(); f['aus']=pd.DataFrame({'SAP':['1','1'],'Cl.':['0191','0195'],'Inicio':['01.02.2026','05.02.2026'],'Final':['10.02.2026','15.02.2026'],'Días':['10,00','11,00']})
        r,_=reconcile(f,date(2026,6,30),True)
        self.assertEqual(r.set_index('SAP').loc['1','Días no remunerados'],15)
    def test_sf_cutoff_and_manual_option(self):
        f=fixtures(); f['sf'].loc[2]=['1','2','01/07/2026','CO_vacatio']; f['sf']['Posting Type']=['Normal','Normal','Ajuste manual']
        r,_=reconcile(f,date(2026,6,30),True)
        self.assertEqual(r.set_index('SAP').loc['1','SF'],5.25)
        r,_=reconcile(f,date(2026,6,30),True,True)
        self.assertEqual(r.set_index('SAP').loc['1','SF'],7.25)
    def test_sf_empty_account_is_not_a_zero_balance(self):
        f=fixtures(); f['sf'].loc[2]=['3','','','CO_vacatio']
        r,audit=reconcile(f,date(2026,6,30),True,include_orphans=True)
        self.assertEqual(audit['Cuentas SF sin movimientos'],1)
        self.assertTrue(pd.isna(r.set_index('SAP').loc['3','SF']))
    def test_reader_encoding_header_and_duplicates(self):
        txt='Título\n------\n|Nº pers.|Cantidad|Cantidad|\n----------------\n|001|15,00|3,00|\n|Nº pers.|Cantidad|Cantidad|\n'
        for encoding in ['cp1252','utf-16','utf-8-sig']:
            r=read_file(txt.encode(encoding),'sap.txt')
            self.assertEqual(list(r.SAP),['001'])
            self.assertIn('Cantidad.1',r)
    def test_current_provision_not_previous_balance(self):
        r,audit=reconcile(fixtures(),date(2026,6,30),True)
        self.assertEqual(r.set_index('SAP').loc['1','Provisión'],6.5)
        self.assertEqual(audit['Campo provisión'],'Dias N Sld')
    def test_approved_future_is_separate_and_does_not_adjust_raw_balances(self):
        f=fixtures()
        f['pending']=pd.DataFrame({'SAP':['1','1'],'approvalStatus':['APPROVED','APPROVED'],'externalCode':['CO_vacatio']*2,'startDate':['01/07/2026']*2,'endDate':['05/07/2026']*2,'quantityInDays':['4','4']})
        r,_=reconcile(f,date(2026,6,30),True)
        r=r.set_index('SAP')
        self.assertEqual(r.loc['1','Futuras aprobadas SF'],4)
        self.assertEqual(r.loc['1','Pendientes SF'],0)
        self.assertEqual(r.loc['1','SF'],5.25)
    def test_only_unused_trailing_csv_columns_may_be_absent(self):
        data=b'User ID,Amount posted,Posting Date,Unused\n1,5,30/06/2026\n'
        r=read_file(data,'sf.csv')
        self.assertEqual(r['Amount posted'].iloc[0],'5')
        with self.assertRaisesRegex(ValueError,'cantidad de columnas'):
            read_file(b'User ID,Amount posted,Posting Date\n1,5\n','sf.csv')
    def test_old_balance_alone_is_not_accepted_as_current(self):
        f=fixtures(); f['prov']=f['prov'].drop(columns='Dias N Sld')
        with self.assertRaisesRegex(ValueError,'Dias N Sld'):
            reconcile(f,date(2026,6,30),True)
    def test_partial_history_does_not_fabricate_independent_balance(self):
        f=fixtures(); f['master']['Fecha']=['01.01.2020','01.01.2026']
        r,_=reconcile(f,date(2026,6,30),True,history_start=date(2025,1,1))
        self.assertTrue(pd.isna(r.set_index('SAP').loc['1','Cálculo']))
        self.assertFalse(pd.isna(r.set_index('SAP').loc['2','Cálculo']))
    def test_only_master_at_cutoff_and_three_states(self):
        f=fixtures(); f['sf'].loc[2]=['3','10','30/06/2026','CO_vacatio']
        f['master'].loc[1,'Fecha']='01.07.2026'
        r,audit=reconcile(f,date(2026,6,30),True)
        self.assertEqual(list(r.SAP),['1'])
        self.assertEqual(audit['Identificadores fuera del maestro'],1)
        self.assertEqual(audit['Ingresos posteriores al corte'],1)
    def test_missing_data_is_separate_from_numerical_discrepancy(self):
        f=fixtures(); f['prov']['Dias N Sld']=['5,25','5,25']
        r,_=reconcile(f,date(2026,6,30),False)
        self.assertEqual(set(r.Estado),{'Sin información'})
    def test_single_sheet_and_formula_injection(self):
        r,_=reconcile(fixtures(),date(2026,6,30),True)
        r.loc[0,'Nombre']='=1+1'
        wb=load_workbook(io.BytesIO(excel_bytes(r)))
        self.assertEqual(wb.sheetnames,['Conciliación'])
        ws=wb.active
        self.assertTrue(ws.auto_filter.ref)
        names=[c for row in ws.iter_rows() for c in row if c.value=='=1+1']
        self.assertTrue(names and all(c.data_type=='s' for c in names))
        self.assertEqual(ws['A2'].value,'Creado por Equipo de automatizaciones nómina')
        self.assertEqual(ws['A6'].value,'Estado')

    def test_auto_history_start_uses_earliest_absence(self):
        f=fixtures()
        r,audit=reconcile(f,date(2026,6,30),True,history_start='auto')
        start=min(pd.to_datetime(f['aus']['Inicio'],dayfirst=True)).date()
        self.assertTrue(audit['Cobertura detectada automáticamente'])
        self.assertEqual(audit['Cobertura declarada desde'],str(start))

if __name__=='__main__': unittest.main()
